"""Automated Verification of Clinical Excel Report Generation
SIH 2026 | PS SIH26113 | Team: BERSERK TECHIES | Team ID: TEAM-147
"""

import io
import unittest
import openpyxl
from backend.simulation.loop import SimulationEngine


class TestClinicalReport(unittest.TestCase):

    def setUp(self):
        self.engine = SimulationEngine(dt=0.01)

    def test_excel_report_generation_and_sheets(self):
        """Verify Excel workbook is created with all 4 sheets and valid data rows."""
        self.engine.start_demo()
        
        # Run 200 steps (2 seconds) of demonstration
        for _ in range(200):
            self.engine.step()

        # Generate in-memory Excel bytes
        excel_bytes = self.engine.generate_excel_report()
        self.assertGreater(len(excel_bytes), 1000, "Excel report size should be substantial")
        self.assertTrue(excel_bytes.startswith(b"PK\x03\x04"), "Must be a valid OpenXML zip archive")

        # Load back into openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        
        self.assertIn("Clinical Summary", sheet_names)
        self.assertIn("Kinematics & Torques", sheet_names)
        self.assertIn("Virtual Sensors", sheet_names)
        self.assertIn("AAN & Safety Events", sheet_names)

        # Inspect Sheet 1
        ws_sum = wb["Clinical Summary"]
        self.assertIn("MoveAssist", str(ws_sum["A1"].value))
        self.assertIn("SIH 2026", str(ws_sum["A2"].value))

        # Inspect Sheet 2 (Kinematics & Torques)
        ws_kin = wb["Kinematics & Torques"]
        self.assertEqual(ws_kin["A1"].value, "Timestamp (s)")
        self.assertEqual(ws_kin["B1"].value, "Kinematic Mode")
        self.assertEqual(ws_kin["F1"].value, "Knee Angle (°)")
        self.assertGreater(ws_kin.max_row, 50, "Should contain multiple time-series records")

        # Inspect Sheet 3 (Virtual Sensors)
        ws_sen = wb["Virtual Sensors"]
        self.assertEqual(ws_sen["A1"].value, "Timestamp (s)")
        self.assertEqual(ws_sen["C1"].value, "Thigh IMU Pitch (°)")
        self.assertEqual(ws_sen["K1"].value, "Encoder Angle (°)")
        self.assertEqual(ws_sen["N1"].value, "FSR Heel (N)")
        self.assertGreater(ws_sen.max_row, 50)

        # Inspect CSV export
        csv_data = self.engine.generate_csv_report()
        self.assertIn("knee_angle_deg", csv_data)
        self.assertIn("tau_req_nm", csv_data)
        self.assertIn("imu_thigh_pitch_deg", csv_data)
        self.assertIn("mech_power_w", csv_data)
        self.assertIn("human_strength_used_pct", csv_data)


if __name__ == "__main__":
    unittest.main()
