"""Clinical Engineering Multi-Sheet Excel & CSV Report Generator
SIH 2026 | PS SIH26113 | Team: BERSERK TECHIES | Team ID: TEAM-147
"""

import io
import datetime
from typing import List, Dict, Any

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from backend.config import SCIENTIFIC_DISCLAIMER


class ClinicalReportGenerator:
    """Generates structured multi-sheet Microsoft Excel (.xlsx) workbooks

    and CSV exports of the 5-phase clinical demonstration trial data.
    """

    @staticmethod
    def generate_excel(trial_records: List[Dict[str, Any]],
                       anthro_summary: Dict[str, Any]) -> bytes:
        """Constructs an Excel workbook (.xlsx) with:

        1. Executive Clinical Summary
        2. Kinematics & Torques (Angles, Velocities, Torques)
        3. Virtual Sensors (IMUs, FSR Ground Reaction, Encoder, sEMG)
        4. AAN Adaptation & Safety Events
        """
        if not OPENPYXL_AVAILABLE:
            raise RuntimeError("openpyxl is not installed.")

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Common styling constants
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sub_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        sub_font = Font(name="Calibri", size=10, bold=True, color="38BDF8")
        title_font = Font(name="Calibri", size=14, bold=True, color="0284C7")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # =====================================================================
        # SHEET 1: EXECUTIVE CLINICAL SUMMARY
        # =====================================================================
        ws_sum = wb.create_sheet(title="Clinical Summary")
        ws_sum.views.sheetView[0].showGridLines = True

        ws_sum["A1"] = "MoveAssist — AI-Assisted Lower-Limb Rehabilitation Exoskeleton"
        ws_sum["A1"].font = title_font
        ws_sum["A2"] = f"SIH 2026 | PS SIH26113 | Team: BERSERK TECHIES (TEAM-147) | Generated: {now_str}"
        ws_sum["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

        ws_sum["A4"] = "SCIENTIFIC DISCLAIMER:"
        ws_sum["A4"].font = bold_font
        ws_sum["B4"] = SCIENTIFIC_DISCLAIMER
        ws_sum["B4"].font = Font(name="Calibri", size=10, italic=True, color="991B1B")

        # Patient Demographics Table
        ws_sum["A6"] = "1. Patient Anthropometrics & Biomechanical Parameters"
        ws_sum["A6"].font = sub_font

        anthro_rows = [
            ("Patient Height (m)", anthro_summary.get("height_m", 1.75)),
            ("Patient Mass (kg)", anthro_summary.get("mass_kg", 72.0)),
            ("Thigh Segment Mass (kg)", anthro_summary.get("thigh", {}).get("mass_kg", 7.20)),
            ("Thigh Length (m)", anthro_summary.get("thigh", {}).get("length_m", 0.43)),
            ("Shank Segment Mass (kg)", anthro_summary.get("shank", {}).get("mass_kg", 3.35)),
            ("Shank Length (m)", anthro_summary.get("shank", {}).get("length_m", 0.43)),
            ("Foot Segment Mass (kg)", anthro_summary.get("foot", {}).get("mass_kg", 1.04)),
            ("Total Distal Mass (kg)", anthro_summary.get("total_distal_mass_kg", 4.39)),
            ("Knee Moment of Inertia (kg·m²)", anthro_summary.get("I_knee_total_kgm2", 0.385)),
            ("Calibration State", anthro_summary.get("calibration_state", "CALIBRATED")),
        ]

        ws_sum["A7"] = "Parameter"
        ws_sum["B7"] = "Value"
        ws_sum["A7"].font = header_font
        ws_sum["A7"].fill = header_fill
        ws_sum["B7"].font = header_font
        ws_sum["B7"].fill = header_fill

        for idx, (param, val) in enumerate(anthro_rows, start=8):
            ws_sum[f"A{idx}"] = param
            ws_sum[f"B{idx}"] = val
            ws_sum[f"A{idx}"].font = regular_font
            ws_sum[f"B{idx}"].font = bold_font
            ws_sum[f"A{idx}"].border = thin_border
            ws_sum[f"B{idx}"].border = thin_border

        # 5-Phase Demonstration Performance Summary Table
        start_demo_row = 20
        ws_sum[f"A{start_demo_row}"] = "2. 5-Phase Clinical Demonstration Evaluation Matrix"
        ws_sum[f"A{start_demo_row}"].font = sub_font

        phase_headers = [
            "Phase #", "Phase Name", "Samples", "Mean Knee Angle (°)",
            "Mean Req Torque (Nm)", "Mean User Torque (Nm)", "Mean Assist Torque (Nm)",
            "Assistance %", "User Strength", "Peak Fatigue", "Safety Mode"
        ]

        row_h = start_demo_row + 1
        for col_idx, col_name in enumerate(phase_headers, start=1):
            cell = ws_sum.cell(row=row_h, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Group trial records by phase
        phase_groups: Dict[int, List[Dict[str, Any]]] = {1: [], 2: [], 3: [], 4: [], 5: []}
        for rec in trial_records:
            p_num = rec.get("demo_phase_num", 0)
            if p_num in phase_groups:
                phase_groups[p_num].append(rec)

        phase_names = {
            1: "Phase 1: Weak Patient",
            2: "Phase 2: Patient Improvement",
            3: "Phase 3: Muscle Fatigue",
            4: "Phase 4: Sudden Deterioration",
            5: "Phase 5: Recovery & Re-adaptation"
        }

        row_p = row_h + 1
        for p_idx in range(1, 6):
            group = phase_groups[p_idx]
            n_samples = len(group)
            if n_samples > 0:
                mean_angle = sum(r["knee_angle_deg"] for r in group) / n_samples
                mean_req = sum(r["tau_req_nm"] for r in group) / n_samples
                mean_user = sum(r["tau_user_nm"] for r in group) / n_samples
                mean_cmd = sum(r["tau_cmd_nm"] for r in group) / n_samples
                mean_assist_pct = sum(r["assistance_pct"] for r in group) / n_samples
                mean_strength = sum(r["user_strength"] for r in group) / n_samples
                peak_fatigue = max(r["fatigue_index"] for r in group)
                mode = group[-1].get("safety_mode", "NORMAL_AAN")
            else:
                mean_angle = mean_req = mean_user = mean_cmd = mean_assist_pct = mean_strength = peak_fatigue = 0.0
                mode = "N/A"

            vals = [
                p_idx, phase_names[p_idx], n_samples,
                round(mean_angle, 2), round(mean_req, 2), round(mean_user, 2),
                round(mean_cmd, 2), f"{round(mean_assist_pct, 1)}%",
                f"{round(mean_strength * 100, 1)}%", round(peak_fatigue, 3), mode
            ]

            for c_idx, v in enumerate(vals, start=1):
                cell = ws_sum.cell(row=row_p, column=c_idx, value=v)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if c_idx in (1, 3, 8, 9, 11) else "right")
            row_p += 1

        # =====================================================================
        # SHEET 2: CONTINUOUS KINEMATICS & TORQUES
        # =====================================================================
        ws_kin = wb.create_sheet(title="Kinematics & Torques")
        ws_kin.views.sheetView[0].showGridLines = True

        kin_headers = [
            "Timestamp (s)", "Kinematic Mode", "Demo Phase", "Gait Phase", "Gait Cycle (%)",
            "Knee Angle (°)", "Thigh Angle (°)", "Shank Angle (°)", "Knee Velocity (°/s)",
            "Required Torque (Nm)", "Gravity Torque (Nm)", "Inertial Torque (Nm)", "External Torque (Nm)",
            "User Estimated Torque (Nm)", "Assistance Deficit (Nm)", "Commanded Assistance (Nm)",
            "Actuator Output Torque (Nm)", "Assistance (%)", "User Contribution (%)",
            "Mechanical Power (W)", "Electrical Power (W)", "Human Strength Used (%)", "Exo Strength Used (%)"
        ]

        for col_idx, h in enumerate(kin_headers, start=1):
            cell = ws_kin.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, r in enumerate(trial_records, start=2):
            row_vals = [
                r["timestamp_s"], r.get("kinematic_mode", "WALK"), r["demo_phase_name"], r["gait_phase"], r["gait_cycle_pct"],
                r["knee_angle_deg"], r["thigh_angle_deg"], r["shank_angle_deg"], r["knee_vel_deg_s"],
                r["tau_req_nm"], r["tau_grav_nm"], r["tau_inert_nm"], r["tau_ext_nm"],
                r["tau_user_nm"], r["tau_deficit_nm"], r["tau_cmd_nm"],
                r["tau_act_nm"], r["assistance_pct"], r["user_pct"],
                r.get("mech_power_w", 0.0), r.get("elec_power_w", 0.0),
                r.get("human_strength_used_pct", 50.0), r.get("exo_strength_used_pct", 50.0)
            ]
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws_kin.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font

        # =====================================================================
        # SHEET 3: VIRTUAL SENSORS STREAM
        # =====================================================================
        ws_sen = wb.create_sheet(title="Virtual Sensors")
        ws_sen.views.sheetView[0].showGridLines = True

        sen_headers = [
            "Timestamp (s)", "Demo Phase",
            "Thigh IMU Pitch (°)", "Thigh IMU Gyro Z (°/s)", "Thigh IMU Accel X (g)", "Thigh IMU Accel Y (g)",
            "Shank IMU Pitch (°)", "Shank IMU Gyro Z (°/s)", "Shank IMU Accel X (g)", "Shank IMU Accel Y (g)",
            "Encoder Angle (°)", "Encoder Velocity (°/s)", "Encoder Counts",
            "FSR Heel (N)", "FSR Metatarsal (N)", "FSR Toe (N)", "Total Vertical GRF (N)", "Stance Detected",
            "sEMG RMS Envelope", "sEMG Raw (uV)", "sEMG Median Freq (Hz)"
        ]

        for col_idx, h in enumerate(sen_headers, start=1):
            cell = ws_sen.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, r in enumerate(trial_records, start=2):
            row_vals = [
                r["timestamp_s"], r["demo_phase_name"],
                r["imu_thigh_pitch_deg"], r["imu_thigh_gyro_deg_s"], r["imu_thigh_accel_x"], r["imu_thigh_accel_y"],
                r["imu_shank_pitch_deg"], r["imu_shank_gyro_deg_s"], r["imu_shank_accel_x"], r["imu_shank_accel_y"],
                r["encoder_angle_deg"], r["encoder_vel_deg_s"], r["encoder_counts"],
                r["fsr_heel_n"], r["fsr_meta_n"], r["fsr_toe_n"], r["fsr_total_grf_n"], r["fsr_is_stance"],
                r["emg_envelope"], r["emg_raw_uV"], r["emg_median_freq_hz"]
            ]
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws_sen.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font

        # =====================================================================
        # SHEET 4: AAN ADAPTATION & SAFETY EVENTS
        # =====================================================================
        ws_safe = wb.create_sheet(title="AAN & Safety Events")
        ws_safe.views.sheetView[0].showGridLines = True

        safe_headers = [
            "Timestamp (s)", "Demo Phase", "K_AAN Gain", "Patient Capacity", "Fatigue Index",
            "Safety Mode", "E-Stop Latched", "Motor Temp (°C)", "Active Safety Alerts"
        ]

        for col_idx, h in enumerate(safe_headers, start=1):
            cell = ws_safe.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, r in enumerate(trial_records, start=2):
            row_vals = [
                r["timestamp_s"], r["demo_phase_name"], r["k_aan"], r["user_strength"], r["fatigue_index"],
                r["safety_mode"], "YES" if r["e_stop_latched"] else "NO", r["motor_temp_c"], r["alerts_str"]
            ]
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws_safe.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font

        # Auto-fit column widths across all sheets
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col[:100]:  # inspect first 100 rows for speed
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Save workbook to in-memory bytes buffer
        out_buf = io.BytesIO()
        wb.save(out_buf)
        return out_buf.getvalue()

    @staticmethod
    def generate_csv(trial_records: List[Dict[str, Any]]) -> str:
        """Fallback CSV generator if user requests raw CSV pipe."""
        if not trial_records:
            return "No data recorded."

        headers = list(trial_records[0].keys())
        lines = [",".join(headers)]
        for rec in trial_records:
            lines.append(",".join(str(rec.get(h, "")) for h in headers))
        return "\n".join(lines)
